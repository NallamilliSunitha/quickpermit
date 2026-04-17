# permissions/views.py

import os
from datetime import date, timedelta

from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponseForbidden
from django.views.decorators.http import require_POST
from django.contrib import messages
from django.conf import settings
from django.core.mail import send_mail
from django.utils import timezone
from django.db import transaction, models

from django.contrib.auth.models import User

from accounts.models import UserProfile
from accounts.views import _send_assigned_email
from .models import PermissionRequest, RequestHistory
from core.utils import create_notification, send_email_if_possible
from django.urls import reverse
from core.utils import create_notification
from core.models import Notification


# ══════════════════════════════════════════════════════════════
# HELPERS
# ══════════════════════════════════════════════════════════════

def _full_name_or_username(u):
    full = (f"{u.first_name} {u.last_name}").strip()
    return full if full else (u.username or "User")


def _get_role(user):
    p = UserProfile.objects.filter(user=user).first()
    return (p.role or "").strip().lower() if p else ""


def _get_dept(user):
    p = UserProfile.objects.filter(user=user).first()
    return (p.department or "").strip() if p else ""


ROLE_FLOW = {
    "student" : ["proctor", "staff", "hod", "dean", "principal"],
    "proctor" : ["hod", "dean", "principal"],
    "staff"   : ["hod", "dean", "principal"],
    "hod"     : ["dean", "principal"],
    "dean"    : ["principal"],
    "principal": [],
}


# ══════════════════════════════════════════════════════════════
# EMAIL HELPERS
# ══════════════════════════════════════════════════════════════

def notify_student(req, event, *, actor=None, to_user=None, extra_note=None):
    """
    Sends email to student for events:
    received, forwarded, approved, rejected, auto_escalated
    """
    request_id    = getattr(req, "request_code", None) or f"REQ-{req.id:06d}"
    title         = (req.title or "").strip() or "Permission Request"
    student_name  = _full_name_or_username(req.student)
    student_email = (req.student.email or "").strip()

    if not student_email:
        print(f"[NOTIFY_STUDENT] No email for student {req.student.username} — skipping")
        return

    authority_name = _full_name_or_username(req.request_to) if req.request_to else "N/A"
    action_by      = _full_name_or_username(actor) if actor else "System"
    date_from      = getattr(req, "from_date", None)
    date_to        = getattr(req, "to_date", None)
    date_range     = f"{date_from} to {date_to}" if (date_from and date_to) else "—"
    status         = (req.status or "").upper()
    current_level  = (req.current_level or "").upper()

    subject_map = {
        "received"      : f"[{request_id}] Permission Request Submitted Successfully",
        "forwarded"     : f"[{request_id}] Permission Request Forwarded",
        "approved"      : f"[{request_id}] Permission Request Approved",
        "rejected"      : f"[{request_id}] Permission Request Rejected",
        "auto_escalated": f"[{request_id}] Permission Request Auto-Escalated",
    }
    subject = subject_map.get(event, f"[{request_id}] Permission Request Update")

    lines = [
        f"Hello {student_name},",
        "",
        "This is an update regarding your permission request.",
        "",
        f"Request ID   : {request_id}",
        f"Title        : {title}",
        f"Date Range   : {date_range}",
        f"Status       : {status}",
        f"Current Level: {current_level}",
        f"Assigned To  : {authority_name}",
    ]

    if event == "received":
        lines += ["", "Your request has been successfully submitted and assigned for review."]
    elif event == "forwarded":
        to_user_name = _full_name_or_username(to_user) if to_user else "Next Authority"
        lines += ["", f"Your request has been forwarded to: {to_user_name}.", f"Forwarded By: {action_by}"]
    elif event == "approved":
        lines += ["", "Your request has been approved.", f"Approved By: {action_by}"]
    elif event == "rejected":
        lines += ["", "Your request has been rejected.", f"Rejected By: {action_by}"]
    elif event == "auto_escalated":
        to_user_name = _full_name_or_username(to_user) if to_user else authority_name
        lines += [
            "",
            "Your request was not acted upon within the specified time and has been auto-escalated.",
            f"Escalated To: {to_user_name}",
        ]

    if extra_note:
        lines += ["", f"Note: {extra_note}"]

    lines += [
        "",
        "You can track your request from your dashboard.",
        "https://sunitha11.pythonanywhere.com/dashboard/",
        "",
        "Regards,",
        "CampusIQ Team"
    ]

    try:
        send_mail(
            subject        = subject,
            message        = "\n".join(lines),
            from_email     = settings.DEFAULT_FROM_EMAIL,
            recipient_list = [student_email],
            fail_silently  = False,
        )
        print(f"[NOTIFY_STUDENT] Email sent to {student_email} for event={event}")
    except Exception as e:
        print(f"[NOTIFY_STUDENT] Email failed: {e}")

    # Parent email if available
    try:
        p = UserProfile.objects.filter(user=req.student).first()
        if p and getattr(p, "parent_email", None):
            send_email_if_possible(
                p.parent_email,
                f"Update on your ward's permission request ({req.request_code})",
                f"Request: {req.title}\nStatus: {req.status.upper()}\nStudent: {req.student.username}"
            )
    except Exception as e:
        print(f"[NOTIFY_STUDENT] Parent email failed: {e}")


def notify_assignee(req, event, *, actor=None, from_user=None, extra_note=None):
    """
    Sends email to assigned authority when:
    event="assigned"  → new request created
    event="forwarded" → request forwarded to this authority
    """
    if not req.request_to:
        return

    to_email = (req.request_to.email or "").strip()
    if not to_email:
        print(f"[NOTIFY_ASSIGNEE] No email for {req.request_to.username} — skipping")
        return

    request_id     = getattr(req, "request_code", None) or f"REQ-{req.id:06d}"
    title          = (req.title or "").strip() or "Permission Request"
    authority_name = _full_name_or_username(req.request_to)
    student_name   = _full_name_or_username(req.student)
    date_from      = getattr(req, "from_date", None)
    date_to        = getattr(req, "to_date", None)
    date_range     = f"{date_from} to {date_to}" if (date_from and date_to) else "—"
    action_by      = _full_name_or_username(actor) if actor else "System"
    from_name      = _full_name_or_username(from_user) if from_user else student_name

    if event == "assigned":
        subject = f"[{request_id}] New Permission Request Assigned to You"
        body_lines = [
            f"Hello {authority_name},",
            "",
            "A new permission request has been assigned to you for review.",
            "",
            f"Request ID : {request_id}",
            f"Student    : {student_name} ({req.student.username})",
            f"Title      : {title}",
            f"Date Range : {date_range}",
            f"Urgent     : {'YES' if req.is_urgent else 'NO'}",
        ]
    else:  # forwarded
        subject = f"[{request_id}] Permission Request Forwarded to You"
        body_lines = [
            f"Hello {authority_name},",
            "",
            "A permission request has been forwarded to you for review.",
            "",
            f"Request ID   : {request_id}",
            f"Student      : {student_name} ({req.student.username})",
            f"Title        : {title}",
            f"Date Range   : {date_range}",
            f"Forwarded By : {action_by}",
            f"From         : {from_name}",
        ]

    if extra_note:
        body_lines += ["", f"Note: {extra_note}"]

    body_lines += [
        "",
        "Please login to CampusIQ and take action.",
        "https://sunitha11.pythonanywhere.com/dashboard/",
        "",
        "Regards,",
        "CampusIQ Team"
    ]

    try:
        send_mail(
            subject        = subject,
            message        = "\n".join(body_lines),
            from_email     = settings.DEFAULT_FROM_EMAIL,
            recipient_list = [to_email],
            fail_silently  = False,
        )
        print(f"[NOTIFY_ASSIGNEE] Email sent to {to_email} for event={event}")
    except Exception as e:
        print(f"[NOTIFY_ASSIGNEE] Email failed: {e}")


# ══════════════════════════════════════════════════════════════
# BASIC VIEWS
# ══════════════════════════════════════════════════════════════

@login_required
def index(request):
    return render(request, "permissions/index.html")


def extract_text_from_uploaded_file(file_field):
    if not file_field:
        return ("", None)
    try:
        path = file_field.path
    except Exception:
        return ("", "File path not available.")
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        try:
            import docx
            doc  = docx.Document(path)
            text = "\n".join([p.text for p in doc.paragraphs]).strip()
            return (text, None if text else "DOCX has no readable text.")
        except Exception as e:
            return ("", f"Could not read DOCX: {e}")
    if ext == ".pdf":
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(path)
            parts  = [page.extract_text() or "" for page in reader.pages]
            text   = "\n\n".join([p.strip() for p in parts if p.strip()])
            if not text:
                return ("", "PDF looks scanned — OCR needed.")
            return (text, None)
        except Exception as e:
            return ("", f"Could not read PDF: {e}")
    if ext == ".doc":
        return ("", "Old .DOC format not supported. Upload .DOCX or PDF.")
    return ("", f"Unsupported file type: {ext}")


@login_required
def view_request(request, id):
    req = get_object_or_404(PermissionRequest, id=id)
    auto_event    = RequestHistory.objects.filter(request=req, action="auto_escalated").order_by("-created_at").first()
    last_reassign = RequestHistory.objects.filter(request=req, action="reassigned").order_by("-created_at").first()
    return render(request, "permissions/view_request.html", {
        "req"           : req,
        "display_reason": req.reason or "",
        "extract_error" : None,
        "last_reassign" : last_reassign,
        "auto_event"    : auto_event,
    })


# ══════════════════════════════════════════════════════════════
# APPROVE / REJECT
# ══════════════════════════════════════════════════════════════

@login_required
def approve_request(request, id):
    req = get_object_or_404(PermissionRequest, id=id)
    req.status = "approved"
    req.save()

    RequestHistory.objects.create(
        request   = req,
        action    = "approved",
        from_role = _get_role(request.user),
        to_role   = None,
        actor     = request.user,
        note      = "Approved"
    )

    create_notification(
    user=req.student,   # requester (student/staff/hod etc)
    title="Permission Approved",
    message=f"Your request {req.request_code} has been approved.",
    link=reverse("request_detail", args=[req.id])
)
    notify_student(req, "approved", actor=request.user)
    return redirect("dashboard")



@login_required
def reject_request(request, id):
    req = get_object_or_404(PermissionRequest, id=id)
    req.status = "rejected"
    req.save()

    RequestHistory.objects.create(
        request   = req,
        action    = "rejected",
        from_role = _get_role(request.user),
        to_role   = None,
        actor     = request.user,
        note      = "Rejected"
    )

    create_notification(
    user=req.student,   # requester (student/staff/hod etc)
    title="Permission Rejected",
    message=f"Your request {req.request_code} has been rejected.",
    link=reverse("request_detail", args=[req.id])
)
    notify_student(req, "rejected", actor=request.user)
    return redirect("dashboard")



@login_required
def forward_request(request, id):
    req = get_object_or_404(PermissionRequest, id=id)
    return redirect("dashboard")


# ══════════════════════════════════════════════════════════════
# FORWARD UI / FORWARD DO
# ══════════════════════════════════════════════════════════════

@login_required
def forward_ui(request, pk):
    req = get_object_or_404(PermissionRequest, pk=pk)

    my_profile = UserProfile.objects.filter(user=request.user).first()
    if not my_profile:
        return JsonResponse({"error": "no profile"}, status=403)

    my_role = (my_profile.role or "").strip().lower()
    if my_role == "student":
        return JsonResponse({"error": "student blocked"}, status=403)
    if req.request_to_id != request.user.id:
        return JsonResponse({"error": "not assigned"}, status=403)

    my_dept       = my_profile.department
    allowed_roles = ROLE_FLOW.get(my_role, [])
    selected_role = (request.GET.get("role") or "").strip().lower()

    users = []
    if selected_role and selected_role in allowed_roles:
        qs = UserProfile.objects.filter(role=selected_role)
    # Only filter by department for staff/proctor/hod level
    # Dean and Principal are institution-wide — no department filter
        if selected_role not in ["dean", "principal"]:
            qs = qs.filter(department=my_dept)
        qs    = qs.select_related("user")
        users = list(qs.values("user__id", "user__username", "user__first_name", "user__last_name"))
    return JsonResponse({"allowed_roles": allowed_roles, "users": users})


@login_required
@require_POST
def forward_do(request, pk):
    req = get_object_or_404(PermissionRequest, pk=pk)

    my_profile = UserProfile.objects.filter(user=request.user).first()
    if not my_profile:
        return JsonResponse({"ok": False, "error": "Profile not found"}, status=403)

    my_role = (my_profile.role or "").strip().lower()
    if my_role == "student":
        return JsonResponse({"ok": False, "error": "Students cannot forward"}, status=403)
    if req.request_to_id != request.user.id:
        return JsonResponse({"ok": False, "error": "Not assigned to you"}, status=403)

    my_dept        = my_profile.department
    allowed_roles  = ROLE_FLOW.get(my_role, [])
    target_role    = (request.POST.get("target_role") or "").strip().lower()
    target_user_id = request.POST.get("target_user_id")
    comment        = (request.POST.get("comment") or "").strip()

    if target_role not in allowed_roles:
        return JsonResponse({"ok": False, "error": "Not allowed role"}, status=403)

    # Dean and Principal are institution-wide — no department filter
    if target_role in ["dean", "principal"]:
        target_profile = UserProfile.objects.filter(
        user_id=target_user_id, role=target_role
    ).select_related("user").first()
    else:
        target_profile = UserProfile.objects.filter(
        user_id=target_user_id, role=target_role, department=my_dept
    ).select_related("user").first()

    if not target_profile:
        return JsonResponse({"ok": False, "error": "User not found"}, status=404)

    new_level = (target_profile.role or "").strip().lower()

    req.request_to    = target_profile.user
    req.status        = "pending"
    req.current_level = new_level
    req.save(update_fields=["request_to", "status", "current_level", "updated_at"])

    RequestHistory.objects.create(
        request   = req,
        action    = "forwarded",
        from_role = my_role,
        to_role   = new_level,
        actor     = request.user,
        note      = comment if comment else "Forwarded"
    )
    create_notification(
    user=target_profile.user,
    title=f"New Request Assigned: {req.request_code}",
    message=(
        f"Request '{req.title}' has been forwarded to you by "
        f"{request.user.username}."
        + (f"\nComment: {comment}" if comment else "")
    ),
    link=f"/permissions/view/{req.id}/"
)

# 🔔 Notify REQUEST CREATOR (student/staff/etc)
    create_notification(
    user=req.student,
    title=f"Your Request Forwarded: {req.request_code}",
    message=(
        f"Your request '{req.title}' has been forwarded "
        f"to {target_profile.user.username} ({new_level.upper()})."
    ),
    link=f"/permissions/track/{req.id}/"
)
    notify_assignee(req, "forwarded", actor=request.user, from_user=request.user, extra_note=comment or None)
    notify_student(req, "forwarded", actor=request.user, to_user=target_profile.user, extra_note=comment or None)

    return JsonResponse({"ok": True})


# ══════════════════════════════════════════════════════════════
# TRACK / DELETE
# ══════════════════════════════════════════════════════════════

@login_required
def track_request(request, id):
    req = get_object_or_404(PermissionRequest, id=id)
    if req.student_id != request.user.id:
        return HttpResponseForbidden("Not allowed")
    history = RequestHistory.objects.filter(request=req).order_by("created_at")
    return render(request, "permissions/track_request.html", {
        "request_obj": req,
        "history"    : history,
    })


@login_required
def delete_request(request, id):
    if request.method != "POST":
        return HttpResponseForbidden("POST only")

    req = get_object_or_404(PermissionRequest, id=id)

    if req.student_id != request.user.id:
        return HttpResponseForbidden("You can delete only your own requests")
    if (req.status or "").strip().lower() != "pending":
        return HttpResponseForbidden("Only pending requests can be deleted")

    profile = UserProfile.objects.filter(user=request.user).first()
    my_role = (profile.role or "").strip().lower() if profile else ""

    RequestHistory.objects.create(
        request   = req,
        action    = "deleted",
        from_role = my_role,
        to_role   = None,
        actor     = request.user,
        note      = "Deleted by requester"
    )
    req.delete()
    messages.success(request, "Request deleted successfully.")
    return redirect("my_requests")


# ══════════════════════════════════════════════════════════════
# BULK FORWARD
# ══════════════════════════════════════════════════════════════

@login_required
@require_POST
def bulk_forward_do(request):
    my_profile = UserProfile.objects.filter(user=request.user).first()
    if not my_profile:
        return JsonResponse({"ok": False, "error": "Profile not found"}, status=403)

    my_role = (my_profile.role or "").strip().lower()
    if my_role == "student":
        return JsonResponse({"ok": False, "error": "Students cannot forward"}, status=403)

    my_dept        = my_profile.department
    allowed_roles  = ROLE_FLOW.get(my_role, [])
    target_role    = (request.POST.get("target_role") or "").strip().lower()
    target_user_id = request.POST.get("target_user_id")
    ids            = [int(x) for x in request.POST.getlist("request_ids") if str(x).isdigit()]

    if not ids:
        return JsonResponse({"ok": False, "error": "No requests selected"}, status=400)
    if target_role not in allowed_roles:
        return JsonResponse({"ok": False, "error": "Not allowed role"}, status=403)

    target_profile = UserProfile.objects.filter(
        user_id=target_user_id, role=target_role, department=my_dept
    ).select_related("user").first()

    if not target_profile:
        return JsonResponse({"ok": False, "error": "Target user not found"}, status=404)

    new_role = (target_profile.role or "").strip().lower()
    qs       = PermissionRequest.objects.filter(id__in=ids, request_to=request.user, status="pending")
    updated  = 0
    skipped  = []

    with transaction.atomic():
        for req in qs.select_related("student"):
            req.request_to    = target_profile.user
            req.current_level = new_role
            req.status        = "pending"
            req.save(update_fields=["request_to", "current_level", "status", "updated_at"])
            notify_assignee(req, "forwarded", actor=request.user, from_user=request.user)
            RequestHistory.objects.create(
                request   = req,
                action    = "forwarded",
                from_role = my_role,
                to_role   = new_role,
                actor     = request.user,
                note      = f"Forwarded to {target_profile.user.username}"
            )
            updated += 1

        valid_ids = set(qs.values_list("id", flat=True))
        skipped   = [rid for rid in ids if rid not in valid_ids]

    return JsonResponse({"ok": True, "updated": updated, "skipped": skipped, "target": target_profile.user.username})


# ══════════════════════════════════════════════════════════════
# REASSIGN UI / REASSIGN DO
# ══════════════════════════════════════════════════════════════

@login_required
def reassign_ui(request, pk):
    req = get_object_or_404(PermissionRequest, pk=pk)

    my_profile = UserProfile.objects.filter(user=request.user).first()
    if not my_profile:
        return JsonResponse({"ok": False, "error": "Profile not found"}, status=403)

    my_role = (my_profile.role or "").strip().lower()
    if req.request_to_id != request.user.id:
        return JsonResponse({"ok": False, "error": "Not assigned to you"}, status=403)

    student_profile = UserProfile.objects.filter(user=req.student).first()
    student_dept    = student_profile.department if student_profile else None

    if my_role == "staff":
        qs = UserProfile.objects.filter(role__in=["staff", "proctor"], department=my_profile.department).exclude(user=request.user).select_related("user")
    elif my_role == "hod":
        qs = UserProfile.objects.filter(role__in=["staff", "proctor"], department=my_profile.department).exclude(user=request.user).select_related("user")
    elif my_role == "dean":
        qs = UserProfile.objects.filter(role__in=["staff", "hod", "proctor"], department=student_dept).exclude(user=request.user).select_related("user")
    elif my_role == "principal":
        qs = UserProfile.objects.filter(
            models.Q(role__in=["staff", "hod", "proctor"], department=student_dept) | models.Q(role__iexact="dean")
        ).exclude(user=request.user).select_related("user")
    else:
        return JsonResponse({"ok": False, "error": "Not allowed"}, status=403)

    users = [{
        "id"        : p.user.id,
        "username"  : p.user.username,
        "first_name": p.user.first_name,
        "last_name" : p.user.last_name,
        "role"      : (p.role or "").strip().lower(),
    } for p in qs]

    return JsonResponse({"ok": True, "users": users})


@login_required
@require_POST
def reassign_do(request, pk):
    req = get_object_or_404(PermissionRequest, pk=pk)

    if req.request_to != request.user:
        return JsonResponse({"ok": False, "error": "Not allowed"}, status=403)

    target_user_id = request.POST.get("target_user_id")
    comment        = (request.POST.get("comment") or "").strip()

    if not target_user_id:
        return JsonResponse({"ok": False, "error": "No user selected"}, status=400)

    target_profile = UserProfile.objects.filter(user_id=target_user_id).select_related("user").first()
    if not target_profile:
        return JsonResponse({"ok": False, "error": "User not found"}, status=404)

    actor_profile = UserProfile.objects.filter(user=request.user).first()
    actor_role    = (actor_profile.role or "").strip().lower() if actor_profile else ""

    if actor_role in ("hod", "dean", "principal") and not comment:
        return JsonResponse({"ok": False, "error": "Comment is required for HOD/Dean/Principal"}, status=400)

    # ── Update request ─────────────────────────────
    old_role          = req.current_level
    req.request_to    = target_profile.user
    req.current_level = (target_profile.role or "").strip().lower()
    req.status        = "pending"
    req.save(update_fields=["request_to", "current_level", "status", "updated_at"])

    # ── History ─────────────────────────────────────
    RequestHistory.objects.create(
        request   = req,
        action    = "reassigned",
        from_role = actor_role,
        to_role   = req.current_level,
        actor     = request.user,
        note      = comment if comment else "Reassigned"
    )

    # ── In-app notification — new assignee ──────────
    create_notification(
        user    = target_profile.user,
        title   = f"Request Reassigned to You: {req.request_code}",
        message = (
            f"Request {req.request_code} ('{req.title}') was reassigned to you "
            f"by {_full_name_or_username(request.user)}."
            + (f"\nComment: {comment}" if comment else "")
        ),
        link = f"/permissions/view/{req.id}/"
    )

    # ── In-app notification — student ───────────────
    create_notification(
        user    = req.student,
        title   = f"Your Request Reassigned: {req.request_code}",
        message = (
            f"Your request '{req.title}' has been reassigned "
            f"from {old_role.upper()} to {req.current_level.upper()}."
        ),
        link = f"/permissions/track/{req.id}/"
    )

    # ── Email to new assignee ───────────────────────
    try:
        if target_profile.user.email:
            assignee_name = _full_name_or_username(target_profile.user)
            student_name  = _full_name_or_username(req.student)
            actor_name    = _full_name_or_username(request.user)
            send_mail(
                subject = f"[{req.request_code}] Permission Request Reassigned to You",
                message = (
                    f"Hello {assignee_name},\n\n"
                    f"A permission request has been reassigned to you.\n\n"
                    f"Request Code : {req.request_code}\n"
                    f"Title        : {req.title}\n"
                    f"Student      : {student_name} ({req.student.username})\n"
                    f"From Date    : {req.from_date}\n"
                    f"To Date      : {req.to_date}\n"
                    f"Reassigned By: {actor_name}\n"
                    + (f"Comment      : {comment}\n" if comment else "")
                    + f"\nPlease login to CampusIQ and take action.\n"
                    f"https://sunitha11.pythonanywhere.com/dashboard/\n\n"
                    f"Regards,\nCampusIQ Team"
                ),
                from_email     = settings.DEFAULT_FROM_EMAIL,
                recipient_list = [target_profile.user.email],
                fail_silently  = False,
            )
            print(f"[REASSIGN] Email sent to new assignee: {target_profile.user.email}")
        else:
            print(f"[REASSIGN] No email for new assignee — skipping")
    except Exception as e:
        print(f"[REASSIGN] New assignee email failed: {e}")

    # ── Email to student ────────────────────────────
    try:
        if req.student.email:
            student_name  = _full_name_or_username(req.student)
            assignee_name = _full_name_or_username(target_profile.user)
            send_mail(
                subject = f"[{req.request_code}] Your Request Has Been Reassigned",
                message = (
                    f"Hello {student_name},\n\n"
                    f"Your permission request has been reassigned to a new authority.\n\n"
                    f"Request Code  : {req.request_code}\n"
                    f"Title         : {req.title}\n"
                    f"From Date     : {req.from_date}\n"
                    f"To Date       : {req.to_date}\n"
                    f"Reassigned To : {assignee_name} ({req.current_level.upper()})\n"
                    + (f"Comment       : {comment}\n" if comment else "")
                    + f"\nYou can track your request from your dashboard.\n"
                    f"https://sunitha11.pythonanywhere.com/dashboard/\n\n"
                    f"Regards,\nCampusIQ Team"
                ),
                from_email     = settings.DEFAULT_FROM_EMAIL,
                recipient_list = [req.student.email],
                fail_silently  = False,
            )
            print(f"[REASSIGN] Student email sent to {req.student.email}")
        else:
            print(f"[REASSIGN] No student email — skipping")
    except Exception as e:
        print(f"[REASSIGN] Student email failed: {e}")

    return JsonResponse({"ok": True})


# ══════════════════════════════════════════════════════════════
# SUGGEST FORWARD TARGETS
# ══════════════════════════════════════════════════════════════

def suggest_forward_targets(current_user, req):
    my_profile = UserProfile.objects.filter(user=current_user).first()
    if not my_profile:
        return []
    my_role       = (my_profile.role or "").strip().lower()
    dept          = my_profile.department
    allowed_roles = ROLE_FLOW.get(my_role, [])
    if not allowed_roles:
        return []
    qs          = UserProfile.objects.filter(department=dept, role__in=allowed_roles).select_related("user")
    suggestions = []
    for p in qs:
        pending_load = PermissionRequest.objects.filter(request_to=p.user, status="pending").count()
        full         = (p.user.get_full_name() or "").strip()
        suggestions.append({
            "id"           : p.user.id,
            "name"         : full if full else p.user.username,
            "role"         : p.role,
            "pending_count": pending_load
        })
    suggestions.sort(key=lambda x: x["pending_count"])
    return suggestions[:5]


@login_required
def forward_suggestions_api(request, pk):
    req = get_object_or_404(PermissionRequest, pk=pk)
    if req.request_to_id != request.user.id:
        return HttpResponseForbidden("Not allowed")
    my_profile = UserProfile.objects.filter(user=request.user).select_related("user").first()
    if not my_profile:
        return JsonResponse({"ok": False, "error": "Profile not found"}, status=403)
    dept = my_profile.department
    qs   = (
        UserProfile.objects
        .filter(department=dept)
        .exclude(user=request.user)
        .select_related("user")
        .values("user__id", "user__username", "user__first_name", "user__last_name", "role")
        .order_by("role", "user__first_name", "user__username")
    )
    users = []
    for u in qs:
        full = (f"{u['user__first_name']} {u['user__last_name']}").strip()
        users.append({
            "id"      : u["user__id"],
            "username": u["user__username"],
            "name"    : full if full else u["user__username"],
            "role"    : (u["role"] or "").lower(),
        })
    return JsonResponse({"ok": True, "request_id": req.id, "department": dept, "users": users})


# ══════════════════════════════════════════════════════════════
# AI INSIGHT
# ══════════════════════════════════════════════════════════════

def _month_range(d):
    if not d:
        d = timezone.localdate()
    month_start = d.replace(day=1)
    if month_start.month == 12:
        next_month_start = date(month_start.year + 1, 1, 1)
    else:
        next_month_start = date(month_start.year, month_start.month + 1, 1)
    return month_start, next_month_start


@login_required
def ai_insight_api(request, pk):
    req = get_object_or_404(PermissionRequest, pk=pk)

    my_profile = UserProfile.objects.filter(user=request.user).first()
    my_role    = (my_profile.role or "").strip().lower() if my_profile else ""

    if my_role == "student":
        return JsonResponse({"ok": False, "error": "Not allowed"}, status=403)
    if req.request_to_id != request.user.id:
        return JsonResponse({"ok": False, "error": "Not assigned to you"}, status=403)

    applicant         = req.student
    applicant_profile = UserProfile.objects.filter(user=applicant).first()

    if req.from_date:
        base_date = req.from_date
    elif req.applied_at:
        applied_at = req.applied_at
        if timezone.is_naive(applied_at):
            applied_at = timezone.make_aware(applied_at)
        base_date = timezone.localtime(applied_at).date()
    else:
        base_date = timezone.now().date()

    month_start, next_month_start = _month_range(base_date)
    leave_type = (req.title or "").strip() or "General"

    month_all_qs  = PermissionRequest.objects.filter(student=applicant, from_date__gte=month_start, from_date__lt=next_month_start)
    month_type_qs = month_all_qs.filter(title=leave_type)

    month_all_taken  = month_all_qs.filter(status="approved").count()
    month_type_taken = month_type_qs.filter(status="approved").count()
    month_all_total  = month_all_qs.count()
    month_type_total = month_type_qs.count()
    month_rejected   = month_all_qs.filter(status="rejected").count()
    all_time_total   = PermissionRequest.objects.filter(student=applicant).count()
    all_time_approved= PermissionRequest.objects.filter(student=applicant, status="approved").count()
    all_time_rejected= PermissionRequest.objects.filter(student=applicant, status="rejected").count()
    approval_rate    = round((all_time_approved / all_time_total * 100)) if all_time_total > 0 else 0

    ai_summary = ""
    ai_recommendation = ""
    ai_score = None
    ai_flags = []

    try:
        from groq import Groq
        client         = Groq(api_key=settings.GROQ_API_KEY)
        applicant_name = f"{applicant.first_name} {applicant.last_name}".strip() or applicant.username
        applicant_role = (applicant_profile.role or "student").title() if applicant_profile else "Student"
        dept           = (applicant_profile.department or "Unknown") if applicant_profile else "Unknown"

        response = client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            max_tokens=500,
            messages=[
                {"role": "system", "content": "You are an AI assistant helping college authorities make fair decisions on permission requests. Be concise, fair, and professional. Do not use emojis."},
                {"role": "user", "content": f"""Analyze this permission request:

Applicant: {applicant_name} ({applicant_role}, {dept})
Request Title: {leave_type}
From: {req.from_date} To: {req.to_date}
Reason: {(req.reason or "Not provided")[:300]}
Urgent: {"Yes" if req.is_urgent else "No"}

This Month: Total={month_all_total}, Approved={month_all_taken}, Rejected={month_rejected}, Same Type={month_type_total}
All Time: Total={all_time_total}, Approval Rate={approval_rate}%

Respond in this exact format:
SCORE: [0-100]
RECOMMENDATION: [APPROVE / REJECT / REVIEW]
FLAGS: [comma separated concerns or NONE]
SUMMARY: [2-3 sentences]"""}
            ]
        )

        text = response.choices[0].message.content.strip()
        for line in text.split("\n"):
            line = line.strip()
            if line.startswith("SCORE:"):
                try: ai_score = int(line.replace("SCORE:", "").strip())
                except: ai_score = None
            elif line.startswith("RECOMMENDATION:"):
                ai_recommendation = line.replace("RECOMMENDATION:", "").strip()
            elif line.startswith("FLAGS:"):
                flags_raw = line.replace("FLAGS:", "").strip()
                ai_flags  = [f.strip() for f in flags_raw.split(",") if f.strip() and f.strip().upper() != "NONE"]
            elif line.startswith("SUMMARY:"):
                ai_summary = line.replace("SUMMARY:", "").strip()

    except Exception as e:
        print(f"[AI Insight ERROR] {e}")
        ai_summary        = "Unable to generate AI insight. Please review manually."
        ai_recommendation = "REVIEW"

    return JsonResponse({
        "ok"              : True,
        "summary"         : ai_summary,
        "score"           : ai_score,
        "recommendation"  : ai_recommendation,
        "flags"           : ai_flags,
        "approval_rate"   : approval_rate,
        "month"           : month_start.strftime("%B %Y"),
        "applicant"       : applicant.username,
        "applicant_name"  : f"{applicant.first_name} {applicant.last_name}".strip() or applicant.username,
        "leave_type"      : leave_type,
        "same_type_taken" : month_type_taken,
        "all_types_taken" : month_all_taken,
        "same_type_total" : month_type_total,
        "all_types_total" : month_all_total,
        "month_rejected"  : month_rejected,
        "all_time_total"  : all_time_total,
        "all_time_approved": all_time_approved,
        "all_time_rejected": all_time_rejected,
    })
# permissions/views.py

from django.http import JsonResponse
from django.conf import settings
from .utils import send_urgent_reminders

def send_reminders_view(request):
    key = request.GET.get("key")

    if key != settings.CRON_SECRET_KEY:
        return JsonResponse({"error": "Unauthorized"}, status=403)

    count = send_urgent_reminders()
    return JsonResponse({"status": "ok", "sent": count})


# ══════════════════════════════════════════════════════════════
# GENERATE PERMISSION LETTER
# ══════════════════════════════════════════════════════════════

@login_required
def generate_permission_letter(request):
    if request.method != "POST":
        return JsonResponse({"ok": False, "error": "Invalid method"})
    try:
        from groq import Groq
        from datetime import date

        title      = request.POST.get("title", "").strip()
        leave_type = request.POST.get("leave_type", "").strip()
        from_date  = request.POST.get("from_date", "").strip()
        to_date    = request.POST.get("to_date", "").strip()
        authority  = request.POST.get("authority", "The Principal").strip()
        details    = request.POST.get("details", "").strip()

        user    = request.user
        name    = f"{user.first_name} {user.last_name}".strip() or user.username
        profile = user.userprofile
        dept    = profile.department or "Computer Science and Engineering"
        roll    = getattr(profile, "roll_number", "") or user.username
        role    = (profile.role or "").strip().lower()

        today = date.today().strftime("%d %B %Y")

        role_map = {
            "student"  : ("Student", f"Roll Number: {roll}", "I assure you that I will complete all missed assignments and academic work upon my return."),
            "staff"    : ("Staff Member", f"Employee ID: {roll}", "I assure you that I will make necessary arrangements for my classes and duties to be covered during my absence."),
            "proctor"  : ("Proctor", f"Employee ID: {roll}", "I assure you that I will ensure all student responsibilities are properly managed during my absence."),
            "hod"      : ("Head of Department", f"Employee ID: {roll}", "I assure you that all departmental responsibilities will be properly delegated and the department will function smoothly during my absence."),
            "dean"     : ("Dean", f"Employee ID: {roll}", "I assure you that all administrative responsibilities will be duly delegated and institutional operations will continue smoothly during my absence."),
            "principal": ("Principal", f"Employee ID: {roll}", "I assure you that all institutional responsibilities will be properly managed and delegated during my absence."),
        }

        role_label, id_line, assurance = role_map.get(role, (role.title(), f"ID: {roll}", "I assure you that all responsibilities will be managed during my absence."))

        client = Groq(api_key=settings.GROQ_API_KEY)

        response = client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            max_tokens=1000,
            messages=[
                {
                    "role": "system",
                    "content": (
                        f"You are an expert at writing formal official letters for Indian engineering colleges. "
                        f"You strictly follow proper Indian letter writing format with From address, To address, "
                        f"date, subject line, salutation, well-structured body paragraphs, and formal closing. "
                        f"The letter is written by a {role_label}. "
                        f"Output ONLY the complete letter. No explanations, no notes, nothing extra before or after the letter."
                    )
                },
                {
                    "role": "user",
                    "content": f"""Write a formal permission/leave letter following this EXACT format:

From,
{name}
{role_label}
Department of {dept}
Aditya Engineering College (Autonomous)
Surampalem, Andhra Pradesh

To,
{authority}
Aditya Engineering College (Autonomous)
Surampalem, Andhra Pradesh

Date: {today}

Subject: [Write a specific, meaningful one-line subject about {title} — {leave_type}]

Respected Sir/Madam,

[Paragraph 1: Introduce yourself — mention your name, role, department, and {id_line}. State the purpose of this letter clearly.]

[Paragraph 2: Explain the reason for leave/permission clearly. Mention exact dates from {from_date} to {to_date}. Reason type is {leave_type}. Additional details: {details if details else "Not provided"}. Be specific and professional.]

[Paragraph 3: Write this exactly — {assurance}]

[Paragraph 4: Request the authority to kindly grant the permission/leave. Express gratitude for their consideration. Mention availability to provide any further information if required.]

Thanking you,
Yours faithfully,

{name}
{role_label}
Department of {dept}
{id_line}
Date: {today}

STRICT RULES YOU MUST FOLLOW:
- Fill ALL details with actual information provided — no placeholders like [Your Name] or [Date]
- Subject line must be specific — example: "Request for Medical Leave from 25th to 27th March 2026"
- Each paragraph must be complete and professional
- Output ONLY the letter — absolutely nothing before From, or after the date at the bottom
- Maintain formal tone throughout appropriate for a {role_label}"""
                }
            ]
        )

        letter = response.choices[0].message.content.strip()
        return JsonResponse({"ok": True, "letter": letter})

    except Exception as e:
        print("Letter generation error:", e)
        return JsonResponse({"ok": False, "error": str(e)})





# ══════════════════════════════════════════════════════════════
# RUN AUTO ESCALATION (via URL key)
# ══════════════════════════════════════════════════════════════

def run_auto_escalation(request):
    secret = request.GET.get("key")
    if secret != settings.AUTO_ESCALATION_KEY:
        return JsonResponse({"ok": False, "error": "Unauthorized"}, status=401)
    from permissions.utils import auto_escalate_permissions
    count = auto_escalate_permissions()
    return JsonResponse({"ok": True, "count": count})
from django.shortcuts import render, get_object_or_404
from .models import PermissionRequest

def request_detail(request, id):
    req = get_object_or_404(PermissionRequest, id=id)

    return render(request, "permissions/request_detail.html", {
        "req": req
    })